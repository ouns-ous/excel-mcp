"""Test suite for Batch 1: Data Quality & Text Intelligence Headless Tools."""

import os
import openpyxl
from excel_engine import ExcelEngine
import server


def test_batch1_data_quality():
    test_dir = os.path.abspath("test_out")
    os.makedirs(test_dir, exist_ok=True)
    fpath = os.path.join(test_dir, "data_quality_test.xlsx")

    # 1. Setup sample workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws["A1"] = "Customer Name"
    ws["B1"] = "Email"
    ws["C1"] = "Phone"
    ws["D1"] = "Status Note"

    ws["A2"] = "   Acme   Corp   "
    ws["B2"] = "john.doe@company.com"
    ws["C2"] = "555-123-4567"
    ws["D2"] = "invoice #INV-9021 sent"

    ws["A3"] = "Acme Corp."
    ws["B3"] = "jane.smith@partner.org"
    ws["C3"] = "555-987-6543"
    ws["D3"] = "paid in full"

    wb.save(fpath)

    # Test 1: Clean whitespace
    ExcelEngine.clean_whitespace_headless(fpath, range_address="A2:A2")
    wb1 = openpyxl.load_workbook(fpath)
    assert wb1["Customers"]["A2"].value == "Acme Corp"
    print("[OK] clean_whitespace_headless passed")

    # Test 2: Case transform
    ExcelEngine.case_transform_headless(fpath, range_address="A2:A2", mode="upper")
    wb2 = openpyxl.load_workbook(fpath)
    assert wb2["Customers"]["A2"].value == "ACME CORP"
    print("[OK] case_transform_headless passed")

    # Test 3: Regex replace
    ExcelEngine.regex_replace_headless(fpath, range_address="D2:D2", pattern=r"INV-(\d+)", replacement=r"REF-\1")
    wb3 = openpyxl.load_workbook(fpath)
    assert "REF-9021" in wb3["Customers"]["D2"].value
    print("[OK] regex_replace_headless passed")

    # Test 4: Extract pattern
    ExcelEngine.extract_by_pattern_headless(fpath, source_col="D", target_col="E", pattern=r"(REF-\d+)", start_row=2)
    wb4 = openpyxl.load_workbook(fpath)
    assert wb4["Customers"]["E2"].value == "REF-9021"
    print("[OK] extract_by_pattern_headless passed")

    # Test 5: Fuzzy deduplicate
    res_dedup = ExcelEngine.fuzzy_deduplicate_headless(fpath, key_column="A", threshold=0.75, action="flag")
    assert res_dedup["duplicates_found"] >= 1
    wb5 = openpyxl.load_workbook(fpath)
    assert wb5["Customers"]["F1"].value == "Fuzzy_Duplicate_Flag"
    print("[OK] fuzzy_deduplicate_headless passed")

    # Test 6: Mask sensitive data
    ExcelEngine.mask_sensitive_data_headless(fpath, range_address="B2:B2", mask_type="email")
    wb6 = openpyxl.load_workbook(fpath)
    assert "@company.com" in wb6["Customers"]["B2"].value
    assert "*" in wb6["Customers"]["B2"].value
    print("[OK] mask_sensitive_data_headless passed")

    # Test 7: Detect column types
    types_rep = ExcelEngine.detect_column_types_headless(fpath)
    assert "Customer Name" in types_rep["columns"]
    print("[OK] detect_column_types_headless passed")

    # Verify tool count on FastMCP
    tool_count = len(server.mcp._tool_manager._tools)
    print(f"Current registered tool count: {tool_count}")
    assert tool_count == 230


if __name__ == "__main__":
    test_batch1_data_quality()
    print("[ALL BATCH 1 TESTS PASSED SUCCESSFULLY!]")

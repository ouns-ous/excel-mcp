"""Test suite for the new batch of headless tools in excel-mcp.
Tests sheet copying, password protection, deletion, and tab colorization.
"""

import os
import openpyxl
from excel_engine import ExcelEngine


def test_mega_headless_tools():
    test_dir = os.path.abspath("test_out")
    os.makedirs(test_dir, exist_ok=True)
    base_file = os.path.join(test_dir, "mega_test.xlsx")

    # 1. Create base file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MainData"
    ws["A1"] = "Metric"
    ws["B1"] = "Value"
    ws["A2"] = "ActiveUsers"
    ws["B2"] = 45000
    wb.save(base_file)

    # 2. Test Copy Sheet
    ExcelEngine.copy_sheet_headless(base_file, "MainData", "MainData_Backup")
    wb_copied = openpyxl.load_workbook(base_file)
    assert "MainData_Backup" in wb_copied.sheetnames
    assert wb_copied["MainData_Backup"]["B2"].value == 45000
    print("[OK] Copy sheet headless passed.")

    # 3. Test Set Sheet Color
    ExcelEngine.set_sheet_color_headless(base_file, "MainData_Backup", "#10B981")
    wb_colored = openpyxl.load_workbook(base_file)
    assert wb_colored["MainData_Backup"].sheet_properties.tabColor.rgb is not None
    print("[OK] Set sheet tab color headless passed.")

    # 4. Test Protect Sheet
    ExcelEngine.protect_sheet_headless(base_file, password="SecretPassword123", sheet_name="MainData")
    wb_prot = openpyxl.load_workbook(base_file)
    assert wb_prot["MainData"].protection.sheet is True
    print("[OK] Protect sheet headless passed.")

    # 5. Test Delete Sheet
    ExcelEngine.delete_sheet_headless(base_file, "MainData_Backup")
    wb_del = openpyxl.load_workbook(base_file)
    assert "MainData_Backup" not in wb_del.sheetnames
    print("[OK] Delete sheet headless passed.")


if __name__ == "__main__":
    test_mega_headless_tools()
    print("[ALL MEGA HEADLESS TESTS PASSED!]")

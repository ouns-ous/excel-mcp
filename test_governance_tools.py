"""Test suite for Batch 3: Document Governance, Data Schema & Multi-Workbook Management."""

import os
import openpyxl
from excel_engine import ExcelEngine
import server


def test_batch3_governance():
    test_dir = os.path.abspath("test_out")
    os.makedirs(test_dir, exist_ok=True)
    fpath = os.path.join(test_dir, "governance_test.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws["A1"] = "Department"
    ws["B1"] = "Employee Name"
    ws["C1"] = "Salary"

    data = [
        ("Engineering", "Alice", 120000),
        ("Engineering", "Bob", 110000),
        ("Marketing", "Carol", 95000),
        ("Finance", "Dave", 105000)
    ]
    for r_idx, row in enumerate(data, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(fpath)

    # 1. Test Audit Log
    ExcelEngine.append_audit_log_headless(
        fpath, event_message="Automated salary adjustment approved", user_name="AuditBot"
    )
    wb_al = openpyxl.load_workbook(fpath)
    assert "_AuditLog" in wb_al.sheetnames
    assert wb_al["_AuditLog"]["C2"].value == "AuditBot"
    print("[OK] append_audit_log_headless passed")

    # 2. Test Batch Export Sheets to CSV
    csv_files = ExcelEngine.batch_export_sheets_to_csv_headless(fpath, output_dir=test_dir)
    assert len(csv_files) >= 2
    assert any("Employees.csv" in f for f in csv_files)
    print("[OK] batch_export_sheets_to_csv_headless passed")

    # 3. Test Split Workbook by Column
    split_files = ExcelEngine.split_workbook_by_column_headless(
        fpath, split_column="A", output_dir=test_dir, sheet_name="Employees"
    )
    assert len(split_files) == 3  # Engineering, Marketing, Finance
    print("[OK] split_workbook_by_column_headless passed")

    # 4. Test Merge Multiple Workbooks
    merged_path = os.path.join(test_dir, "merged_test.xlsx")
    ExcelEngine.merge_multiple_workbooks_headless(split_files, output_path=merged_path, combine_mode="sheets")
    wb_m = openpyxl.load_workbook(merged_path)
    assert len(wb_m.sheetnames) >= 3
    print("[OK] merge_multiple_workbooks_headless passed")

    # 5. Test Validate Schema
    res_val = ExcelEngine.validate_schema_headless(
        fpath, expected_columns=["Department", "Employee Name", "Salary"], required_types={"Salary": "number"}, sheet_name="Employees"
    )
    assert res_val["is_valid"] is True
    print("[OK] validate_schema_headless passed")

    # 6. Test Generate Markdown Data Dictionary
    md_dict = ExcelEngine.generate_markdown_data_dictionary_headless(fpath, sheet_name="Employees")
    assert "# Data Dictionary: Employees" in md_dict
    assert "Salary" in md_dict
    print("[OK] generate_markdown_data_dictionary_headless passed")

    # Verify final tool count
    total_tools = len(server.mcp._tool_manager._tools)
    print(f"Final total registered tool count: {total_tools}")
    assert total_tools == 243


if __name__ == "__main__":
    test_batch3_governance()
    print("[ALL BATCH 3 GOVERNANCE TESTS PASSED!]")

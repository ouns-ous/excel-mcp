"""Test suite for the latest batch of advanced tools in excel-mcp.
Tests headless comparison, hyperlinks extraction, and sheet unhiding.
"""

import os
import openpyxl
from excel_engine import ExcelEngine


def test_compare_and_hyperlinks():
    test_dir = os.path.abspath("test_out")
    os.makedirs(test_dir, exist_ok=True)

    file1 = os.path.join(test_dir, "cmp1.xlsx")
    file2 = os.path.join(test_dir, "cmp2.xlsx")

    # Create file 1
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Sales"
    ws1["A1"] = "Product"
    ws1["B1"] = "Revenue"
    ws1["A2"] = "Laptop"
    ws1["B2"] = 1000
    c_link = ws1["A3"]
    c_link.value = "Docs"
    c_link.hyperlink = "https://example.com/docs"
    wb1.save(file1)

    # Create file 2 with a difference
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Sales"
    ws2["A1"] = "Product"
    ws2["B1"] = "Revenue"
    ws2["A2"] = "Laptop"
    ws2["B2"] = 1500  # Difference here
    ws_hidden = wb2.create_sheet("HiddenSheet")
    ws_hidden.sheet_state = "hidden"
    ws_hidden["A1"] = "Secret"
    wb2.save(file2)

    # 1. Compare workbooks
    diffs = ExcelEngine.compare_two_workbooks(file1, file2)
    assert len(diffs["only_in_file_2"]) == 1
    assert "HiddenSheet" in diffs["only_in_file_2"]
    assert any(d["cell"] == "B2" for d in diffs["cell_differences"])
    print("[OK] Workbook comparison test passed.")

    # 2. Extract hyperlinks
    links = ExcelEngine.extract_hyperlinks(file1)
    assert len(links) >= 1
    assert links[0]["target"] == "https://example.com/docs"
    print("[OK] Hyperlink extraction test passed.")

    # 3. Unhide all sheets
    unhidden_file = os.path.join(test_dir, "unhidden.xlsx")
    ExcelEngine.unhide_all_sheets_headless(file2, unhidden_file)
    wb_unhidden = openpyxl.load_workbook(unhidden_file)
    assert wb_unhidden["HiddenSheet"].sheet_state == "visible"
    print("[OK] Unhide all sheets test passed.")


if __name__ == "__main__":
    test_compare_and_hyperlinks()
    print("[ALL NEW TOOLS TESTS PASSED SUCCESSFULLY!]")

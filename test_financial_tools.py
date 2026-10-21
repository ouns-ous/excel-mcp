"""Test suite for Batch 2: Financial Modeling & Statistical Analysis Headless Tools."""

import os
import openpyxl
from excel_engine import ExcelEngine
import server


def test_batch2_financial_tools():
    test_dir = os.path.abspath("test_out")
    os.makedirs(test_dir, exist_ok=True)
    fpath = os.path.join(test_dir, "financial_test.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Advertising"
    ws["B1"] = "Sales"
    sample_data = [
        (1000, 15000),
        (2000, 28000),
        (3000, 42000),
        (4000, 55000),
        (5000, 71000)
    ]
    for r_idx, (adv, sls) in enumerate(sample_data, start=2):
        ws[f"A{r_idx}"] = adv
        ws[f"B{r_idx}"] = sls
    wb.save(fpath)

    # 1. Test Amortization
    ExcelEngine.add_amortization_schedule_headless(
        fpath, principal=100000, annual_interest_rate=0.05, term_years=5, sheet_name="LoanSchedule"
    )
    wb_am = openpyxl.load_workbook(fpath)
    assert "LoanSchedule" in wb_am.sheetnames
    assert wb_am["LoanSchedule"]["C2"].value > 0
    print("[OK] add_amortization_schedule_headless passed")

    # 2. Test CAGR
    res_cagr = ExcelEngine.add_cagr_analysis_headless(
        fpath, start_value=100000, end_value=250000, periods=5, output_cell="D2", sheet_name="Data"
    )
    assert "cagr_percent" in res_cagr
    wb_cg = openpyxl.load_workbook(fpath)
    assert "CAGR METRIC ANALYSIS" in wb_cg["Data"]["D2"].value
    print("[OK] add_cagr_analysis_headless passed")

    # 3. Test Depreciation
    ExcelEngine.add_depreciation_schedule_headless(
        fpath, asset_cost=50000, salvage_value=5000, useful_life_years=5, sheet_name="MachineryDepr"
    )
    wb_dp = openpyxl.load_workbook(fpath)
    assert "MachineryDepr" in wb_dp.sheetnames
    assert wb_dp["MachineryDepr"]["C2"].value == 9000.0  # (50000-5000)/5
    print("[OK] add_depreciation_schedule_headless passed")

    # 4. Test Correlation Matrix
    res_corr = ExcelEngine.generate_correlation_matrix_headless(
        fpath, columns=["A", "B"], output_cell="G2", sheet_name="Data"
    )
    assert "correlation_matrix" in res_corr
    assert res_corr["correlation_matrix"]["Advertising"]["Sales"] > 0.95
    print("[OK] generate_correlation_matrix_headless passed")

    # 5. Test Summary Statistics
    res_stats = ExcelEngine.add_summary_statistics_headless(
        fpath, column="B", output_cell="J2", sheet_name="Data"
    )
    assert res_stats["statistics"]["Count"] == 5
    assert res_stats["statistics"]["Mean (Average)"] == 42200.0
    print("[OK] add_summary_statistics_headless passed")

    # 6. Test Break Even
    res_be = ExcelEngine.calculate_break_even_headless(
        fpath, fixed_costs=50000, variable_cost_per_unit=20, price_per_unit=50, output_cell="M2", sheet_name="Data"
    )
    assert res_be["break_even_units"] == 1667.0
    print("[OK] calculate_break_even_headless passed")

    # 7. Test Waterfall Breakdown
    ExcelEngine.add_waterfall_breakdown_headless(
        fpath,
        categories=["Start ARR", "New Sales", "Churn", "Upgrades"],
        amounts=[100000, 30000, -10000, 15000],
        output_cell="P2",
        sheet_name="Data"
    )
    wb_wf = openpyxl.load_workbook(fpath)
    assert wb_wf["Data"]["P2"].value == "Step / Category"
    assert wb_wf["Data"]["S6"].value == 135000.0
    print("[OK] add_waterfall_breakdown_headless passed")

    # Check registered tool count
    total_tools = len(server.mcp._tool_manager._tools)
    print(f"Current registered tool count: {total_tools}")
    assert total_tools == 237


if __name__ == "__main__":
    test_batch2_financial_tools()
    print("[ALL BATCH 2 FINANCIAL & STATISTICAL TESTS PASSED!]")

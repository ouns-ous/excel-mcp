"""Headless Excel OOXML Engine using openpyxl.
Allows fast headless creation, reading, and transformation of Excel (.xlsx) files
without requiring Microsoft Excel desktop to be installed or running.
"""

import os
import re
from typing import Any, Dict, List, Optional, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelEngine:
    """Headless processor for .xlsx files using openpyxl."""

    @staticmethod
    def create_workbook(
        output_path: str,
        data: Optional[List[List[Any]]] = None,
        sheet_name: str = "Sheet1",
        header_bg: str = "1F4E79",
        header_fg: str = "FFFFFF"
    ) -> str:
        """Create a new formatted .xlsx workbook headlessly."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        if data:
            for r_idx, row in enumerate(data, start=1):
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    if r_idx == 1:
                        cell.font = Font(bold=True, color=header_fg.lstrip("#"), name="Segoe UI")
                        cell.fill = PatternFill(start_color=header_bg.lstrip("#"), end_color=header_bg.lstrip("#"), fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.font = Font(name="Segoe UI")

            # Auto-fit columns
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        abs_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(abs_path), exist_ok=True)
        wb.save(abs_path)
        return abs_path

    @staticmethod
    def read_sheet_data(
        file_path: str,
        sheet_name: Optional[str] = None
    ) -> Dict[str, Any]:
        """Read all rows and columns from an .xlsx file headlessly."""
        abs_path = os.path.abspath(file_path)
        if not os.path.exists(abs_path):
            raise FileNotFoundError(f"File not found: {abs_path}")

        wb = openpyxl.load_workbook(abs_path, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(list(row))

        return {
            "file": abs_path,
            "sheet": ws.title,
            "total_rows": len(rows),
            "total_cols": len(rows[0]) if rows else 0,
            "rows": rows
        }

    @staticmethod
    def markdown_to_excel(markdown_table: str, output_path: str, sheet_name: str = "Data") -> str:
        """Convert a Markdown formatted table into a styled Excel (.xlsx) spreadsheet."""
        lines = [line.strip() for line in markdown_table.strip().splitlines() if line.strip()]
        data = []

        for line in lines:
            if not line.startswith("|") or not line.endswith("|"):
                continue
            cells = [c.strip() for c in line.split("|")[1:-1]]
            # Skip separator line (e.g. |---|---|)
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            data.append(cells)

        if not data:
            raise ValueError("No valid Markdown table detected.")

        return ExcelEngine.create_workbook(output_path, data, sheet_name=sheet_name)

    @staticmethod
    def excel_to_markdown(file_path: str, sheet_name: Optional[str] = None) -> str:
        """Extract table from .xlsx and format as a standard Markdown table."""
        res = ExcelEngine.read_sheet_data(file_path, sheet_name)
        rows = res["rows"]
        if not rows:
            return "*(Empty table)*"

        md_lines = []
        header = [str(c or "") for c in rows[0]]
        md_lines.append("| " + " | ".join(header) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")

        for row in rows[1:]:
            cells = [str(c or "") for c in row]
            # Match header width
            while len(cells) < len(header):
                cells.append("")
            md_lines.append("| " + " | ".join(cells[:len(header)]) + " |")

        return "\n".join(md_lines)

    @staticmethod
    def json_to_excel(json_data: List[Dict[str, Any]], output_path: str, sheet_name: str = "Data") -> str:
        """Convert a list of dictionaries (JSON records) to a formatted Excel file."""
        if not json_data:
            raise ValueError("json_data is empty.")

        headers = list(json_data[0].keys())
        rows = [headers]
        for item in json_data:
            rows.append([item.get(h, "") for h in headers])

        return ExcelEngine.create_workbook(output_path, rows, sheet_name=sheet_name)

    @staticmethod
    def csv_to_excel(csv_path: str, output_path: str, delimiter: str = ",", sheet_name: str = "Data") -> str:
        """Convert a CSV file into a styled Excel (.xlsx) spreadsheet."""
        import csv
        abs_csv = os.path.abspath(csv_path)
        if not os.path.exists(abs_csv):
            raise FileNotFoundError(f"CSV file not found: {abs_csv}")

        rows = []
        with open(abs_csv, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f, delimiter=delimiter)
            for r in reader:
                rows.append(r)

        return ExcelEngine.create_workbook(output_path, rows, sheet_name=sheet_name)

    @staticmethod
    def excel_to_csv(file_path: str, output_csv_path: str, sheet_name: Optional[str] = None, delimiter: str = ",") -> str:
        """Extract a sheet from an Excel file into a CSV file."""
        import csv
        res = ExcelEngine.read_sheet_data(file_path, sheet_name)
        rows = res["rows"]
        abs_out = os.path.abspath(output_csv_path)
        os.makedirs(os.path.dirname(abs_out), exist_ok=True)

        with open(abs_out, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=delimiter)
            for r in rows:
                writer.writerow(r)

        return abs_out

    @staticmethod
    def merge_workbooks(file_paths: List[str], output_path: str) -> str:
        """Merge sheets from multiple .xlsx workbooks into a single workbook."""
        merged_wb = openpyxl.Workbook()
        # Remove default sheet
        merged_wb.remove(merged_wb.active)

        for path in file_paths:
            abs_p = os.path.abspath(path)
            if not os.path.exists(abs_p):
                continue
            src_wb = openpyxl.load_workbook(abs_p, data_only=True)
            base_name = os.path.splitext(os.path.basename(abs_p))[0][:15]

            for s_name in src_wb.sheetnames:
                src_sheet = src_wb[s_name]
                new_title = f"{base_name}_{s_name}"[:31]
                target_sheet = merged_wb.create_sheet(title=new_title)

                for row in src_sheet.iter_rows(values_only=True):
                    target_sheet.append(list(row))

        abs_out = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(abs_out), exist_ok=True)
        merged_wb.save(abs_out)
        return abs_out

    @staticmethod
    def audit_workbook(file_path: str) -> Dict[str, Any]:
        """Deep inspect a workbook headlessly: sheets, dimensions, formulas, file size."""
        abs_p = os.path.abspath(file_path)
        if not os.path.exists(abs_p):
            raise FileNotFoundError(f"File not found: {abs_p}")

        size_kb = round(os.path.getsize(abs_p) / 1024, 2)
        wb = openpyxl.load_workbook(abs_p, data_only=False)

        sheets_meta = []
        total_cells = 0
        total_formulas = 0

        for name in wb.sheetnames:
            ws = wb[name]
            max_r = ws.max_row or 0
            max_c = ws.max_column or 0
            cell_count = max_r * max_c
            total_cells += cell_count

            sheet_formulas = 0
            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    if cell.value and str(cell.value).startswith("="):
                        sheet_formulas += 1
            total_formulas += sheet_formulas

            sheets_meta.append({
                "sheet": name,
                "max_row": max_r,
                "max_column": max_c,
                "total_cells": cell_count,
                "formulas_count": sheet_formulas
            })

        return {
            "file": abs_p,
            "size_kb": size_kb,
            "sheets_count": len(wb.sheetnames),
            "sheets": sheets_meta,
            "total_cells_estimate": total_cells,
            "total_formulas": total_formulas
        }

    @staticmethod
    def batch_replace(file_path: str, replacements: Dict[str, str], output_path: str) -> str:
        """Batch replace placeholders (e.g. {{COMPANY}}) across all sheets in a workbook."""
        abs_in = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(abs_in)

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        val = cell.value
                        for k, v in replacements.items():
                            if k in val:
                                val = val.replace(k, str(v))
                        cell.value = val

        abs_out = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(abs_out), exist_ok=True)
        wb.save(abs_out)
        return abs_out

    @staticmethod
    def extract_column(file_path: str, column: Union[str, int], sheet_name: Optional[str] = None) -> List[Any]:
        """Extract all values from a specific column as a list."""
        abs_in = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(abs_in, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        col_idx = column if isinstance(column, int) else openpyxl.utils.column_index_from_string(column)
        values = []
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            values.append(row[0])
        return values

    @staticmethod
    def extract_metadata(file_path: str) -> Dict[str, Any]:
        """Extract core OOXML document metadata (Author, Created, Title, Modified)."""
        abs_in = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(abs_in)
        p = wb.properties
        return {
            "title": p.title,
            "creator": p.creator,
            "subject": p.subject,
            "description": p.description,
            "created": str(p.created) if p.created else None,
            "modified": str(p.modified) if p.modified else None,
            "lastModifiedBy": p.lastModifiedBy
        }

    @staticmethod
    def set_metadata(file_path: str, title: Optional[str] = None, author: Optional[str] = None, subject: Optional[str] = None, output_path: Optional[str] = None) -> str:
        """Set or update document properties/metadata."""
        abs_in = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(abs_in)
        if title is not None:
            wb.properties.title = title
        if author is not None:
            wb.properties.creator = author
            wb.properties.lastModifiedBy = author
        if subject is not None:
            wb.properties.subject = subject

        dest = os.path.abspath(output_path) if output_path else abs_in
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def compare_two_workbooks(file_path_1: str, file_path_2: str) -> Dict[str, Any]:
        """Compare two .xlsx workbooks headlessly and return structural and cell value differences."""
        f1 = os.path.abspath(file_path_1)
        f2 = os.path.abspath(file_path_2)
        wb1 = openpyxl.load_workbook(f1, data_only=True)
        wb2 = openpyxl.load_workbook(f2, data_only=True)

        sheets1 = set(wb1.sheetnames)
        sheets2 = set(wb2.sheetnames)

        diffs = {
            "file_1": f1,
            "file_2": f2,
            "only_in_file_1": list(sheets1 - sheets2),
            "only_in_file_2": list(sheets2 - sheets1),
            "cell_differences": []
        }

        common_sheets = sheets1.intersection(sheets2)
        for s in common_sheets:
            ws1 = wb1[s]
            ws2 = wb2[s]
            max_r = max(ws1.max_row or 0, ws2.max_row or 0)
            max_c = max(ws1.max_column or 0, ws2.max_column or 0)

            for r in range(1, max_r + 1):
                for c in range(1, max_c + 1):
                    v1 = ws1.cell(row=r, column=c).value
                    v2 = ws2.cell(row=r, column=c).value
                    if v1 != v2:
                        diffs["cell_differences"].append({
                            "sheet": s,
                            "cell": f"{get_column_letter(c)}{r}",
                            "file_1_val": str(v1) if v1 is not None else None,
                            "file_2_val": str(v2) if v2 is not None else None
                        })
                        if len(diffs["cell_differences"]) >= 200:
                            diffs["truncated"] = True
                            return diffs

        return diffs

    @staticmethod
    def extract_hyperlinks(file_path: str, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
        """Extract all embedded URLs and hyperlinks from an .xlsx workbook headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        sheets = [wb[sheet_name]] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets

        links = []
        for ws in sheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.hyperlink and cell.hyperlink.target:
                        links.append({
                            "sheet": ws.title,
                            "cell": cell.coordinate,
                            "text": str(cell.value) if cell.value else "",
                            "target": str(cell.hyperlink.target)
                        })
        return links

    @staticmethod
    def unhide_all_sheets_headless(file_path: str, output_path: Optional[str] = None) -> str:
        """Unhide all hidden and very hidden worksheets headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        for ws in wb.worksheets:
            ws.sheet_state = "visible"

        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def copy_sheet_headless(file_path: str, source_sheet: str, new_sheet_name: str, output_path: Optional[str] = None) -> str:
        """Duplicate a worksheet with all its data, styling, and formulas headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        if source_sheet not in wb.sheetnames:
            raise ValueError(f"Sheet '{source_sheet}' not found in workbook.")
        src = wb[source_sheet]
        target = wb.copy_worksheet(src)
        target.title = new_sheet_name
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def protect_sheet_headless(file_path: str, password: str = "", sheet_name: Optional[str] = None, output_path: Optional[str] = None) -> str:
        """Enable password protection on a worksheet headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        if password:
            ws.protection.set_password(password)
        ws.protection.enable()
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def delete_sheet_headless(file_path: str, sheet_name: str, output_path: Optional[str] = None) -> str:
        """Delete a worksheet from a workbook headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        del wb[sheet_name]
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def set_sheet_color_headless(file_path: str, sheet_name: str, hex_color: str, output_path: Optional[str] = None) -> str:
        """Set worksheet tab color headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        ws = wb[sheet_name]
        ws.sheet_properties.tabColor = hex_color.lstrip("#")
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def get_defined_names_headless(file_path: str) -> List[Dict[str, Any]]:
        """Read all workbook defined names / named ranges headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        names = []
        try:
            for dn in wb.defined_names.values():
                dests = []
                try:
                    dests = [f"{s}:{coord}" for s, coord in dn.destinations]
                except Exception:
                    pass
                names.append({
                    "name": dn.name,
                    "value": str(dn.value),
                    "destinations": dests
                })
        except Exception:
            pass
        return names


    @staticmethod
    def insert_rows_headless(
        file_path: str,
        row_idx: int,
        amount: int = 1,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Insert blank rows headlessly at a specific row position."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.insert_rows(row_idx, amount=amount)
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def delete_rows_headless(
        file_path: str,
        row_idx: int,
        amount: int = 1,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Delete rows headlessly at a specific row position."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.delete_rows(row_idx, amount=amount)
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def insert_cols_headless(
        file_path: str,
        col_idx: int,
        amount: int = 1,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Insert blank columns headlessly at a specific column position."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.insert_cols(col_idx, amount=amount)
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def delete_cols_headless(
        file_path: str,
        col_idx: int,
        amount: int = 1,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Delete columns headlessly at a specific column position."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.delete_cols(col_idx, amount=amount)
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def set_gridlines_headless(file_path: str, show: bool = True, sheet_name: Optional[str] = None, output_path: Optional[str] = None) -> str:
        """Toggle gridlines visibility headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.sheet_view.showGridLines = show
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def set_zoom_headless(file_path: str, zoom_scale: int = 100, sheet_name: Optional[str] = None, output_path: Optional[str] = None) -> str:
        """Set worksheet zoom scale percentage headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.sheet_view.zoomScale = zoom_scale
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def add_image_headless(
        file_path: str,
        image_path: str,
        cell: str = "A1",
        width: Optional[float] = None,
        height: Optional[float] = None,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Embed an image into an Excel spreadsheet headlessly."""
        import openpyxl.drawing.image
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        img = openpyxl.drawing.image.Image(os.path.abspath(image_path))
        if width:
            img.width = width
        if height:
            img.height = height
        ws.add_image(img, cell)
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def auto_fit_columns_headless(file_path: str, sheet_name: Optional[str] = None, output_path: Optional[str] = None) -> str:
        """Compute and auto-fit column widths to text content headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for c in col:
                if c.value is not None:
                    max_len = max(max_len, len(str(c.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def freeze_panes_headless(file_path: str, cell: str = "A2", sheet_name: Optional[str] = None, output_path: Optional[str] = None) -> str:
        """Freeze panes at a specific cell coordinate headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        ws.freeze_panes = cell
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def clean_whitespace_headless(
        file_path: str,
        range_address: Optional[str] = None,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Strip leading/trailing whitespace, collapse internal multiple spaces, and remove control characters from cells."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        cells_to_process = ws[range_address] if range_address else ws.iter_rows()
        if range_address and not isinstance(cells_to_process, tuple):
            cells_to_process = ((cells_to_process,),)
            
        for row in cells_to_process:
            for cell in (row if isinstance(row, (tuple, list)) else (row,)):
                if isinstance(cell.value, str):
                    cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', cell.value)
                    cleaned = re.sub(r'[ \t]+', ' ', cleaned).strip()
                    cell.value = cleaned
                    
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def case_transform_headless(
        file_path: str,
        range_address: str,
        mode: str = "upper",
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Transform text case (upper, lower, title, capitalize) across a range headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        cells_to_process = ws[range_address]
        if not isinstance(cells_to_process, tuple):
            cells_to_process = ((cells_to_process,),)
            
        m = mode.lower().strip()
        for row in cells_to_process:
            for cell in (row if isinstance(row, (tuple, list)) else (row,)):
                if isinstance(cell.value, str):
                    if m == "upper":
                        cell.value = cell.value.upper()
                    elif m == "lower":
                        cell.value = cell.value.lower()
                    elif m == "title":
                        cell.value = cell.value.title()
                    elif m == "capitalize":
                        cell.value = cell.value.capitalize()
                        
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def regex_replace_headless(
        file_path: str,
        range_address: str,
        pattern: str,
        replacement: str,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Find and replace text using regular expressions across a cell range headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        compiled_regex = re.compile(pattern)
        cells_to_process = ws[range_address]
        if not isinstance(cells_to_process, tuple):
            cells_to_process = ((cells_to_process,),)
            
        for row in cells_to_process:
            for cell in (row if isinstance(row, (tuple, list)) else (row,)):
                if isinstance(cell.value, str):
                    cell.value = compiled_regex.sub(replacement, cell.value)
                    
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def extract_by_pattern_headless(
        file_path: str,
        source_col: str,
        target_col: str,
        pattern: str,
        start_row: int = 2,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Extract matching regex patterns (or group 1) from source column into target column."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        compiled_regex = re.compile(pattern)
        max_r = ws.max_row
        for r in range(start_row, max_r + 1):
            s_val = ws[f"{source_col}{r}"].value
            if s_val is not None:
                match = compiled_regex.search(str(s_val))
                if match:
                    extracted = match.group(1) if match.groups() else match.group(0)
                    ws[f"{target_col}{r}"].value = extracted
                else:
                    ws[f"{target_col}{r}"].value = None
                    
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def fuzzy_deduplicate_headless(
        file_path: str,
        key_column: str = "A",
        threshold: float = 0.85,
        action: str = "flag",
        start_row: int = 2,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> dict:
        """Detect approximate duplicate rows based on string similarity in a key column."""
        from difflib import SequenceMatcher
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        max_r = ws.max_row
        rows_data = []
        for r in range(start_row, max_r + 1):
            val = str(ws[f"{key_column}{r}"].value or "").strip().lower()
            rows_data.append((r, val))
            
        duplicate_pairs = []
        seen_duplicates = set()
        for i in range(len(rows_data)):
            r1, val1 = rows_data[i]
            if not val1 or r1 in seen_duplicates:
                continue
            for j in range(i + 1, len(rows_data)):
                r2, val2 = rows_data[j]
                if not val2 or r2 in seen_duplicates:
                    continue
                ratio = SequenceMatcher(None, val1, val2).ratio()
                if ratio >= threshold:
                    duplicate_pairs.append({"row": r2, "matched_with_row": r1, "score": round(ratio, 3)})
                    seen_duplicates.add(r2)
                    
        if action.lower() == "flag":
            flag_col = get_column_letter(ws.max_column + 1)
            ws[f"{flag_col}1"].value = "Fuzzy_Duplicate_Flag"
            ws[f"{flag_col}1"].font = Font(bold=True)
            for r in range(start_row, max_r + 1):
                ws[f"{flag_col}{r}"].value = "DUPLICATE" if r in seen_duplicates else "UNIQUE"
        elif action.lower() == "remove":
            for r in sorted(list(seen_duplicates), reverse=True):
                ws.delete_rows(r)
                
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return {
            "file": dest,
            "total_rows_scanned": len(rows_data),
            "duplicates_found": len(duplicate_pairs),
            "matches": duplicate_pairs[:20]
        }

    @staticmethod
    def mask_sensitive_data_headless(
        file_path: str,
        range_address: str,
        mask_type: str = "email",
        mask_char: str = "*",
        visible_chars: int = 4,
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Mask sensitive data like emails, credit cards, or phone numbers in a cell range headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        cells_to_process = ws[range_address]
        if not isinstance(cells_to_process, tuple):
            cells_to_process = ((cells_to_process,),)
            
        def _mask_val(text: str) -> str:
            if not text:
                return text
            mt = mask_type.lower()
            if mt == "email" and "@" in text:
                user, domain = text.split("@", 1)
                masked_user = user[0] + (mask_char * max(len(user) - 2, 1)) + (user[-1] if len(user) > 1 else "")
                return f"{masked_user}@{domain}"
            elif mt in ("credit_card", "phone", "ssn"):
                clean_digits = re.sub(r'\D', '', text)
                if len(clean_digits) > visible_chars:
                    masked = (mask_char * (len(clean_digits) - visible_chars)) + clean_digits[-visible_chars:]
                    return masked
                return mask_char * len(text)
            else:
                if len(text) > visible_chars:
                    return (mask_char * (len(text) - visible_chars)) + text[-visible_chars:]
                return mask_char * len(text)

        for row in cells_to_process:
            for cell in (row if isinstance(row, (tuple, list)) else (row,)):
                if cell.value is not None:
                    cell.value = _mask_val(str(cell.value))
                    
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def detect_column_types_headless(
        file_path: str,
        sheet_name: Optional[str] = None,
        sample_rows: int = 100
    ) -> dict:
        """Infer semantic data types (integer, float, date, boolean, email, text) across columns."""
        import datetime
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        cols_summary = {}
        max_c = ws.max_column
        max_r = min(ws.max_row, sample_rows + 1)
        
        for c in range(1, max_c + 1):
            col_letter = get_column_letter(c)
            header = ws.cell(row=1, column=c).value or f"Col_{col_letter}"
            types_count = {"integer": 0, "float": 0, "date": 0, "boolean": 0, "email": 0, "text": 0}
            null_count = 0
            samples = []
            
            for r in range(2, max_r + 1):
                v = ws.cell(row=r, column=c).value
                if v is None or v == "":
                    null_count += 1
                    continue
                if len(samples) < 3:
                    samples.append(str(v))
                if isinstance(v, bool):
                    types_count["boolean"] += 1
                elif isinstance(v, int):
                    types_count["integer"] += 1
                elif isinstance(v, float):
                    types_count["float"] += 1
                elif isinstance(v, (datetime.date, datetime.datetime)):
                    types_count["date"] += 1
                else:
                    s = str(v).strip()
                    if re.match(r"^[^@]+@[^@]+\.[^@]+$", s):
                        types_count["email"] += 1
                    elif re.match(r"^-?\d+$", s):
                        types_count["integer"] += 1
                    elif re.match(r"^-?\d*\.\d+$", s):
                        types_count["float"] += 1
                    else:
                        types_count["text"] += 1
                        
            inferred = max(types_count, key=types_count.get) if any(types_count.values()) else "empty"
            cols_summary[str(header)] = {
                "column_letter": col_letter,
                "inferred_type": inferred,
                "null_count": null_count,
                "sample_values": samples
            }
            
        return {
            "sheet": ws.title,
            "total_columns": max_c,
            "rows_sampled": max_r - 1,
            "columns": cols_summary
        }

    @staticmethod
    def add_amortization_schedule_headless(
        file_path: str,
        principal: float,
        annual_interest_rate: float,
        term_years: int,
        start_month: str = "2026-01",
        sheet_name: str = "Amortization",
        output_path: Optional[str] = None
    ) -> str:
        """Generate a complete loan amortization schedule table headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb.create_sheet(sheet_name) if sheet_name not in wb.sheetnames else wb[sheet_name]
        
        headers = ["Period", "Payment Date", "Payment ($)", "Principal ($)", "Interest ($)", "Remaining Balance ($)"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c_idx, value=h)
            cell.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
            cell.fill = PatternFill(fill_type="solid", start_color="1F4E79", end_color="1F4E79")
            cell.alignment = Alignment(horizontal="center")
            
        r = annual_interest_rate / 12.0
        n = term_years * 12
        if r > 0:
            monthly_payment = principal * (r * (1 + r)**n) / ((1 + r)**n - 1)
        else:
            monthly_payment = principal / n
            
        balance = principal
        year, month = map(int, start_month.split("-"))
        
        for p in range(1, n + 1):
            interest = balance * r
            principal_paid = monthly_payment - interest
            balance -= principal_paid
            if balance < 0:
                balance = 0.0
                
            date_str = f"{year}-{month:02d}"
            row_vals = [p, date_str, round(monthly_payment, 2), round(principal_paid, 2), round(interest, 2), round(balance, 2)]
            row_idx = p + 1
            for c_idx, val in enumerate(row_vals, start=1):
                c = ws.cell(row=row_idx, column=c_idx, value=val)
                c.font = Font(name="Segoe UI")
                if c_idx >= 3:
                    c.number_format = "$#,##0.00"
                if p % 2 == 0:
                    c.fill = PatternFill(fill_type="solid", start_color="F2F5F9", end_color="F2F5F9")
                    
            month += 1
            if month > 12:
                month = 1
                year += 1
                
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = 18
            
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def add_cagr_analysis_headless(
        file_path: str,
        start_value: float,
        end_value: float,
        periods: int,
        output_cell: str = "B2",
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> dict:
        """Calculate Compound Annual Growth Rate (CAGR) and build formatted summary card headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        cagr = (end_value / start_value) ** (1.0 / periods) - 1.0 if start_value > 0 and periods > 0 else 0.0
        cagr_pct = round(cagr * 100, 2)
        
        from openpyxl.utils.cell import coordinate_to_tuple
        start_r, start_c = coordinate_to_tuple(output_cell)
        
        labels = [
            ("CAGR METRIC ANALYSIS", ""),
            ("Initial Value", start_value),
            ("Final Value", end_value),
            ("Periods (Years)", periods),
            ("CAGR (%)", f"{cagr_pct}%")
        ]
        
        for idx, (lbl, val) in enumerate(labels):
            r = start_r + idx
            c_lbl = ws.cell(row=r, column=start_c, value=lbl)
            c_val = ws.cell(row=r, column=start_c + 1, value=val)
            c_lbl.font = Font(bold=True if idx in (0, 4) else False, name="Segoe UI")
            c_val.font = Font(bold=True if idx == 4 else False, name="Segoe UI")
            if idx == 0:
                c_lbl.fill = PatternFill(fill_type="solid", start_color="0D9488", end_color="0D9488")
                c_lbl.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
            elif idx == 4:
                c_val.fill = PatternFill(fill_type="solid", start_color="CCFBF1", end_color="CCFBF1")
                
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return {
            "cagr_decimal": round(cagr, 4),
            "cagr_percent": f"{cagr_pct}%",
            "start_value": start_value,
            "end_value": end_value,
            "periods": periods,
            "file": dest
        }

    @staticmethod
    def add_depreciation_schedule_headless(
        file_path: str,
        asset_cost: float,
        salvage_value: float,
        useful_life_years: int,
        method: str = "straight_line",
        sheet_name: str = "Depreciation",
        output_path: Optional[str] = None
    ) -> str:
        """Generate a complete multi-year asset depreciation schedule table headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb.create_sheet(sheet_name) if sheet_name not in wb.sheetnames else wb[sheet_name]
        
        headers = ["Year", "Beginning Book Value", "Depreciation Expense", "Accumulated Depreciation", "Ending Book Value"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c_idx, value=h)
            cell.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
            cell.fill = PatternFill(fill_type="solid", start_color="334155", end_color="334155")
            cell.alignment = Alignment(horizontal="center")
            
        depreciable_base = asset_cost - salvage_value
        annual_sl = depreciable_base / useful_life_years if useful_life_years > 0 else 0.0
        
        beg_val = asset_cost
        accum = 0.0
        
        for yr in range(1, useful_life_years + 1):
            if method.lower() == "declining_balance":
                rate = 2.0 / useful_life_years
                expense = min(beg_val * rate, beg_val - salvage_value)
            else:
                expense = annual_sl
                
            accum += expense
            end_val = beg_val - expense
            row_idx = yr + 1
            vals = [yr, round(beg_val, 2), round(expense, 2), round(accum, 2), round(end_val, 2)]
            for c_idx, v in enumerate(vals, start=1):
                c = ws.cell(row=row_idx, column=c_idx, value=v)
                c.font = Font(name="Segoe UI")
                if c_idx >= 2:
                    c.number_format = "$#,##0.00"
                if yr % 2 == 0:
                    c.fill = PatternFill(fill_type="solid", start_color="F8FAFC", end_color="F8FAFC")
            beg_val = end_val
            
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = 22
            
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def generate_correlation_matrix_headless(
        file_path: str,
        columns: List[str],
        start_row: int = 2,
        end_row: Optional[int] = None,
        output_cell: str = "G2",
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> dict:
        """Compute Pearson correlation matrix across numerical columns headlessly."""
        import math
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        last_row = end_row or ws.max_row
        data_series = {}
        headers = []
        for col_letter in columns:
            hdr = str(ws[f"{col_letter}1"].value or col_letter)
            headers.append(hdr)
            vals = []
            for r in range(start_row, last_row + 1):
                v = ws[f"{col_letter}{r}"].value
                try:
                    vals.append(float(v))
                except (ValueError, TypeError):
                    vals.append(0.0)
            data_series[hdr] = vals
            
        def _pearson(x, y):
            n = len(x)
            if n == 0:
                return 1.0
            mx, my = sum(x) / n, sum(y) / n
            num = sum((a - mx) * (b - my) for a, b in zip(x, y))
            den = math.sqrt(sum((a - mx)**2 for a in x) * sum((b - my)**2 for b in y))
            return round(num / den, 3) if den != 0 else 0.0

        matrix = {}
        for h1 in headers:
            matrix[h1] = {}
            for h2 in headers:
                matrix[h1][h2] = _pearson(data_series[h1], data_series[h2])
                
        # Write to sheet
        from openpyxl.utils.cell import coordinate_to_tuple
        start_r, start_c = coordinate_to_tuple(output_cell)
        
        # Header row
        for c_idx, h in enumerate(headers, start=start_c + 1):
            c = ws.cell(row=start_r, column=c_idx, value=h)
            c.font = Font(bold=True, name="Segoe UI")
            c.alignment = Alignment(horizontal="center")
            
        for r_idx, h1 in enumerate(headers, start=start_r + 1):
            c_lbl = ws.cell(row=r_idx, column=start_c, value=h1)
            c_lbl.font = Font(bold=True, name="Segoe UI")
            for c_idx, h2 in enumerate(headers, start=start_c + 1):
                val = matrix[h1][h2]
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(name="Segoe UI")
                cell.number_format = "0.00"
                # Color code
                if val >= 0.7:
                    cell.fill = PatternFill(fill_type="solid", start_color="DCFCE7", end_color="DCFCE7")
                elif val <= -0.7:
                    cell.fill = PatternFill(fill_type="solid", start_color="FEE2E2", end_color="FEE2E2")
                    
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return {"correlation_matrix": matrix, "file": dest}

    @staticmethod
    def add_summary_statistics_headless(
        file_path: str,
        column: str = "B",
        start_row: int = 2,
        end_row: Optional[int] = None,
        output_cell: str = "D2",
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> dict:
        """Compute comprehensive statistical metrics (Count, Mean, Median, Min, Max, StdDev) headlessly."""
        import math
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        last_row = end_row or ws.max_row
        vals = []
        for r in range(start_row, last_row + 1):
            v = ws[f"{column}{r}"].value
            try:
                vals.append(float(v))
            except (ValueError, TypeError):
                continue
                
        n = len(vals)
        if n == 0:
            return {"error": "No numeric data found in column."}
            
        vals.sort()
        mean = sum(vals) / n
        median = (vals[n//2] if n % 2 != 0 else (vals[n//2 - 1] + vals[n//2]) / 2.0)
        variance = sum((x - mean)**2 for x in vals) / (n - 1) if n > 1 else 0.0
        std_dev = math.sqrt(variance)
        
        stats = [
            ("STATISTICAL SUMMARY", f"Column {column}"),
            ("Count", n),
            ("Sum", round(sum(vals), 2)),
            ("Mean (Average)", round(mean, 2)),
            ("Median", round(median, 2)),
            ("Minimum", round(vals[0], 2)),
            ("Maximum", round(vals[-1], 2)),
            ("Range", round(vals[-1] - vals[0], 2)),
            ("Std Deviation", round(std_dev, 2)),
            ("Variance", round(variance, 2))
        ]
        
        from openpyxl.utils.cell import coordinate_to_tuple
        start_r, start_c = coordinate_to_tuple(output_cell)
        
        for idx, (lbl, val) in enumerate(stats):
            r = start_r + idx
            c_lbl = ws.cell(row=r, column=start_c, value=lbl)
            c_val = ws.cell(row=r, column=start_c + 1, value=val)
            c_lbl.font = Font(bold=True if idx == 0 else False, name="Segoe UI")
            c_val.font = Font(name="Segoe UI")
            if idx == 0:
                c_lbl.fill = PatternFill(fill_type="solid", start_color="4F46E5", end_color="4F46E5")
                c_lbl.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
                c_val.fill = PatternFill(fill_type="solid", start_color="4F46E5", end_color="4F46E5")
                c_val.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
            elif idx % 2 == 1:
                c_lbl.fill = PatternFill(fill_type="solid", start_color="F8FAFC", end_color="F8FAFC")
                c_val.fill = PatternFill(fill_type="solid", start_color="F8FAFC", end_color="F8FAFC")
                
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return {"statistics": dict(stats[1:]), "file": dest}

    @staticmethod
    def calculate_break_even_headless(
        file_path: str,
        fixed_costs: float,
        variable_cost_per_unit: float,
        price_per_unit: float,
        target_profit: float = 0.0,
        output_cell: str = "B2",
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> dict:
        """Compute Break-Even units and revenue, contribution margin, and formatted card headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        cm_unit = price_per_unit - variable_cost_per_unit
        cm_ratio = (cm_unit / price_per_unit) if price_per_unit > 0 else 0.0
        be_units = ((fixed_costs + target_profit) / cm_unit) if cm_unit > 0 else 0.0
        be_revenue = be_units * price_per_unit
        
        metrics = [
            ("BREAK-EVEN ANALYSIS", ""),
            ("Fixed Costs", fixed_costs),
            ("Unit Selling Price", price_per_unit),
            ("Unit Variable Cost", variable_cost_per_unit),
            ("Contribution Margin / Unit", round(cm_unit, 2)),
            ("Contribution Margin Ratio", f"{round(cm_ratio * 100, 1)}%"),
            ("Break-Even Volume (Units)", round(be_units, 0)),
            ("Break-Even Revenue ($)", round(be_revenue, 2))
        ]
        
        from openpyxl.utils.cell import coordinate_to_tuple
        start_r, start_c = coordinate_to_tuple(output_cell)
        
        for idx, (lbl, val) in enumerate(metrics):
            r = start_r + idx
            c_lbl = ws.cell(row=r, column=start_c, value=lbl)
            c_val = ws.cell(row=r, column=start_c + 1, value=val)
            c_lbl.font = Font(bold=True if idx in (0, 6, 7) else False, name="Segoe UI")
            c_val.font = Font(bold=True if idx in (6, 7) else False, name="Segoe UI")
            if idx == 0:
                c_lbl.fill = PatternFill(fill_type="solid", start_color="EA580C", end_color="EA580C")
                c_lbl.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
            elif idx in (6, 7):
                c_val.fill = PatternFill(fill_type="solid", start_color="FFEDD5", end_color="FFEDD5")
                
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return {
            "break_even_units": round(be_units, 0),
            "break_even_revenue": round(be_revenue, 2),
            "contribution_margin_unit": round(cm_unit, 2),
            "contribution_margin_ratio": round(cm_ratio, 3),
            "file": dest
        }

    @staticmethod
    def add_waterfall_breakdown_headless(
        file_path: str,
        categories: List[str],
        amounts: List[float],
        output_cell: str = "B2",
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> str:
        """Create structured waterfall breakdown table (Base, Increase, Decrease, Cumulative) headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        from openpyxl.utils.cell import coordinate_to_tuple
        start_r, start_c = coordinate_to_tuple(output_cell)
        
        headers = ["Step / Category", "Amount", "Change Type", "Cumulative Balance"]
        for c_idx, h in enumerate(headers, start=start_c):
            c = ws.cell(row=start_r, column=c_idx, value=h)
            c.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
            c.fill = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
            
        running_total = 0.0
        for idx, (cat, amt) in enumerate(zip(categories, amounts), start=1):
            r = start_r + idx
            running_total += amt
            change_type = "Starting" if idx == 1 else ("Increase" if amt > 0 else "Decrease")
            
            c0 = ws.cell(row=r, column=start_c, value=cat)
            c1 = ws.cell(row=r, column=start_c + 1, value=amt)
            c2 = ws.cell(row=r, column=start_c + 2, value=change_type)
            c3 = ws.cell(row=r, column=start_c + 3, value=round(running_total, 2))
            
            c1.number_format = "$#,##0.00"
            c3.number_format = "$#,##0.00"
            if amt < 0:
                c1.font = Font(color="DC2626", bold=True, name="Segoe UI")
            elif amt > 0 and idx > 1:
                c1.font = Font(color="16A34A", bold=True, name="Segoe UI")
                
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def append_audit_log_headless(
        file_path: str,
        event_message: str,
        user_name: str = "AI Agent",
        log_sheet_name: str = "_AuditLog",
        output_path: Optional[str] = None
    ) -> str:
        """Append an audit trail event log entry to a workbook headlessly."""
        import datetime
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f)
        
        if log_sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(log_sheet_name)
            headers = ["Log ID", "Timestamp (UTC)", "User / Actor", "Action / Event"]
            for c_idx, h in enumerate(headers, start=1):
                c = ws.cell(row=1, column=c_idx, value=h)
                c.font = Font(bold=True, color="FFFFFF", name="Segoe UI")
                c.fill = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
            ws.column_dimensions["A"].width = 12
            ws.column_dimensions["B"].width = 24
            ws.column_dimensions["C"].width = 20
            ws.column_dimensions["D"].width = 50
        else:
            ws = wb[log_sheet_name]
            
        next_row = ws.max_row + 1
        log_id = next_row - 1
        ts = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        
        vals = [log_id, ts, user_name, event_message]
        for c_idx, val in enumerate(vals, start=1):
            c = ws.cell(row=next_row, column=c_idx, value=val)
            c.font = Font(name="Segoe UI")
            if next_row % 2 == 1:
                c.fill = PatternFill(fill_type="solid", start_color="F8FAFC", end_color="F8FAFC")
                
        dest = os.path.abspath(output_path) if output_path else f
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        wb.save(dest)
        return dest

    @staticmethod
    def batch_export_sheets_to_csv_headless(
        file_path: str,
        output_dir: Optional[str] = None,
        delimiter: str = ","
    ) -> List[str]:
        """Export each sheet in an Excel workbook to individual CSV files headlessly."""
        import csv
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f, data_only=True)
        base_dir = os.path.abspath(output_dir) if output_dir else os.path.dirname(f)
        os.makedirs(base_dir, exist_ok=True)
        base_name = os.path.splitext(os.path.basename(f))[0]
        
        exported_files = []
        for name in wb.sheetnames:
            ws = wb[name]
            csv_path = os.path.join(base_dir, f"{base_name}_{name}.csv")
            with open(csv_path, "w", newline="", encoding="utf-8") as csvfile:
                writer = csv.writer(csvfile, delimiter=delimiter)
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        writer.writerow(["" if v is None else v for v in row])
            exported_files.append(csv_path)
        return exported_files

    @staticmethod
    def merge_multiple_workbooks_headless(
        file_paths: List[str],
        output_path: str,
        combine_mode: str = "sheets"
    ) -> str:
        """Merge multiple Excel workbooks into one (by adding sheets or stacking rows) headlessly."""
        dest = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        
        target_wb = openpyxl.Workbook()
        default_sheet = target_wb.active
        
        if combine_mode.lower() == "rows":
            target_ws = default_sheet
            target_ws.title = "Merged_Data"
            header_written = False
            curr_row = 1
            for p in file_paths:
                src_wb = openpyxl.load_workbook(os.path.abspath(p), data_only=True)
                src_ws = src_wb.active
                for r_idx, row_vals in enumerate(src_ws.iter_rows(values_only=True), start=1):
                    if not any(v is not None for v in row_vals):
                        continue
                    if r_idx == 1:
                        if not header_written:
                            for c_idx, val in enumerate(row_vals, start=1):
                                c = target_ws.cell(row=curr_row, column=c_idx, value=val)
                                c.font = Font(bold=True)
                            curr_row += 1
                            header_written = True
                    else:
                        for c_idx, val in enumerate(row_vals, start=1):
                            target_ws.cell(row=curr_row, column=c_idx, value=val)
                        curr_row += 1
        else:
            first = True
            for p in file_paths:
                src_wb = openpyxl.load_workbook(os.path.abspath(p), data_only=True)
                base_name = os.path.splitext(os.path.basename(p))[0]
                for s_name in src_wb.sheetnames:
                    src_ws = src_wb[s_name]
                    target_name = s_name if len(file_paths) == 1 else f"{base_name}_{s_name}"[:31]
                    if first:
                        target_ws = default_sheet
                        target_ws.title = target_name
                        first = False
                    else:
                        target_ws = target_wb.create_sheet(title=target_name)
                    for r_idx, row_vals in enumerate(src_ws.iter_rows(values_only=True), start=1):
                        for c_idx, val in enumerate(row_vals, start=1):
                            target_ws.cell(row=r_idx, column=c_idx, value=val)
                            
        target_wb.save(dest)
        return dest

    @staticmethod
    def split_workbook_by_column_headless(
        file_path: str,
        split_column: str = "A",
        output_dir: Optional[str] = None,
        sheet_name: Optional[str] = None
    ) -> List[str]:
        """Split a worksheet into multiple Excel files partitioned by unique values in a column headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        base_dir = os.path.abspath(output_dir) if output_dir else os.path.dirname(f)
        os.makedirs(base_dir, exist_ok=True)
        
        from openpyxl.utils.cell import column_index_from_string
        col_idx = column_index_from_string(split_column)
        
        headers = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        partitions = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v is not None for v in row):
                continue
            key = str(row[col_idx - 1] if col_idx - 1 < len(row) and row[col_idx - 1] is not None else "Unknown").strip()
            # Clean key for filesystem
            clean_key = re.sub(r'[^a-zA-Z0-9_-]', '_', key)
            if clean_key not in partitions:
                partitions[clean_key] = []
            partitions[clean_key].append(row)
            
        generated_files = []
        for key, rows in partitions.items():
            part_wb = openpyxl.Workbook()
            part_ws = part_wb.active
            part_ws.title = key[:31]
            for c_idx, h in enumerate(headers, start=1):
                c = part_ws.cell(row=1, column=c_idx, value=h)
                c.font = Font(bold=True)
            for r_idx, r_data in enumerate(rows, start=2):
                for c_idx, val in enumerate(r_data, start=1):
                    part_ws.cell(row=r_idx, column=c_idx, value=val)
            part_path = os.path.join(base_dir, f"split_{key}.xlsx")
            part_wb.save(part_path)
            generated_files.append(part_path)
            
        return generated_files

    @staticmethod
    def validate_schema_headless(
        file_path: str,
        expected_columns: List[str],
        required_types: Optional[Dict[str, str]] = None,
        sheet_name: Optional[str] = None
    ) -> dict:
        """Validate worksheet headers and cell types against an expected schema headlessly."""
        f = os.path.abspath(file_path)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        actual_headers = [str(c) if c is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        missing_columns = [col for col in expected_columns if col not in actual_headers]
        
        type_mismatches = []
        if required_types:
            header_to_idx = {h: idx for idx, h in enumerate(actual_headers)}
            for col_name, expected_t in required_types.items():
                if col_name in header_to_idx:
                    c_idx = header_to_idx[col_name]
                    for r_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=min(ws.max_row, 50), values_only=True), start=2):
                        val = row[c_idx] if c_idx < len(row) else None
                        if val is None:
                            continue
                        if expected_t.lower() == "number" and not isinstance(val, (int, float)):
                            try:
                                float(val)
                            except ValueError:
                                type_mismatches.append({"row": r_idx, "column": col_name, "value": str(val), "expected": "number"})
                                
        is_valid = len(missing_columns) == 0 and len(type_mismatches) == 0
        return {
            "is_valid": is_valid,
            "actual_columns": actual_headers,
            "missing_columns": missing_columns,
            "type_mismatches_sample": type_mismatches[:10],
            "file": f
        }

    @staticmethod
    def generate_markdown_data_dictionary_headless(
        file_path: str,
        sheet_name: Optional[str] = None
    ) -> str:
        """Generate a complete Markdown data dictionary catalog from an Excel worksheet headlessly."""
        rep = ExcelEngine.detect_column_types_headless(file_path=file_path, sheet_name=sheet_name)
        cols = rep.get("columns", {})
        
        lines = [
            f"# Data Dictionary: {rep.get('sheet', 'Sheet')}",
            f"**File**: `{os.path.basename(file_path)}` | **Sampled Rows**: {rep.get('rows_sampled', 0)}",
            "",
            "| Column Letter | Field Name | Inferred Data Type | Null Count | Sample Values |",
            "| :--- | :--- | :--- | :--- | :--- |"
        ]
        
        for name, info in cols.items():
            samples = ", ".join([f"`{s}`" for s in info.get("sample_values", [])])
            lines.append(f"| {info.get('column_letter')} | **{name}** | `{info.get('inferred_type')}` | {info.get('null_count')} | {samples} |")
            
        return "\n".join(lines)

